from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook


def parse_workbook(workbook_path: str) -> dict:
    path = Path(workbook_path)
    wb = load_workbook(path, data_only=False)

    named_ranges = [name for name in wb.defined_names.keys()]
    parsed = {"workbook": path.name, "named_ranges": named_ranges, "sheets": []}

    for ws in wb.worksheets:
        formulas = []
        validations = []
        conditional_formats = []

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({"cell": cell.coordinate, "formula": cell.value})

        for dv in ws.data_validations.dataValidation:
            validations.append({"type": dv.type, "formula1": dv.formula1, "sqref": str(dv.sqref)})

        for cf_range in ws.conditional_formatting:
            conditional_formats.append(str(cf_range))

        parsed["sheets"].append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "formulas": formulas,
                "data_validations": validations,
                "conditional_formatting": conditional_formats,
            }
        )

    return parsed
